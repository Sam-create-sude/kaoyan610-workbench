r"""把 610 考研资料包里的 PDF / DOCX 提取成纯文本。

只读原始资料，不改动 E:\610 考研 下的任何文件。
输出到 data/extracted/ 与 data/manifest.json。
同时扫描 data/source_extra/（从 zip 里解压出来的资料）。
"""
import os
import json
import re
from pypdf import PdfReader
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE_DIRS = [
    r"E:/610 考研/资料包/背诵资料(1)/背诵资料",
    os.path.join(HERE, "data", "source_extra"),
]
OUT_DIR = os.path.join(HERE, "data", "extracted")
MANIFEST = os.path.join(HERE, "data", "manifest.json")

COMMON = "的是在有和与设计这个我们不了一个人上中大为"
MIN_CHARS = 100

TOPIC_RULES = [
    ("碳中和", "碳中和与碳达峰"),
    ("碳达峰", "碳中和与碳达峰"),
    ("社会责任", "设计师的社会责任"),
    ("AIGC", "数智与AIGC"),
    ("数智", "数智与AIGC"),
    ("人工智能", "数智与AIGC"),
    ("教育", "设计教育"),
    ("教学", "设计教育"),
    ("数媒", "数媒与数字艺术"),
    ("数字媒体", "数媒与数字艺术"),
    ("视觉作品", "数媒与数字艺术"),
    ("科技", "设计与科技"),
    ("新技术", "设计与科技"),
    ("新材料", "设计与科技"),
    ("环境", "设计与环境"),
    ("生态", "设计与环境"),
    ("可持续", "设计与环境"),
    ("景观", "设计与环境"),
    ("消费", "设计与消费"),
    ("品牌", "设计与消费"),
    ("创新", "设计与创新"),
    ("情感化", "情感化设计"),
    ("人性化", "人性化设计"),
    ("人本", "人性化设计"),
    ("管理", "设计与管理"),
    ("文化", "设计与文化"),
    ("非遗", "设计与文化"),
    ("文创", "设计与文化"),
    ("史论", "设计史论"),
    ("设计学概论", "设计史论"),
    ("名人", "设计史论"),
    ("巴巴纳克", "设计史论"),
    ("万能", "写作模板"),
    ("热点", "设计热点"),
]


def guess_topic(rel_path):
    for kw, topic in TOPIC_RULES:
        if kw in rel_path:
            return topic
    return "综合"


def read_docx(path):
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                lines.append(line)
    return "\n".join(lines)


def read_pdf(path):
    reader = PdfReader(path)
    out = []
    for page in reader.pages:
        t = page.extract_text() or ""
        if t.strip():
            out.append(t.strip())
    return "\n".join(out)


def is_garbled(text):
    if len(text) < 30:
        return True
    hits = sum(text.count(c) for c in COMMON)
    return (hits / len(text)) < 0.01


def clean_text(text):
    text = text.replace("\u3000", " ").replace("\xa0", " ")
    lines = []
    for line in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", line).strip()
        if not line:
            continue
        if re.fullmatch(r"[\d\s\-—–\.、]+", line):
            continue
        if len(line) <= 2:
            continue
        lines.append(line)
    seen = set()
    uniq = []
    for line in lines:
        if line in seen:
            continue
        seen.add(line)
        uniq.append(line)
    return "\n".join(uniq)


def safe_name(rel):
    name = rel.replace("\\", "__").replace("/", "__")
    name = re.sub(r"[^\w\u4e00-\u9fff\.\-]", "_", name)
    return name[:80]


def sig(text, n=300):
    """内容指纹：去掉空白和标点后取前 n 个字符，用来识别重复文件。"""
    return re.sub(r"[\s\W_]", "", text)[:n]


os.makedirs(OUT_DIR, exist_ok=True)

files = []
for base in SOURCE_DIRS:
    if not os.path.isdir(base):
        continue
    for root, dirs, names in os.walk(base):
        for n in names:
            if n.lower().endswith((".pdf", ".docx")):
                path = os.path.join(root, n)
                files.append((path, os.path.relpath(path, base)))
files.sort(key=lambda x: x[1])

manifest = []
skipped = []
seen_names = {}
for i, (path, rel) in enumerate(files, 1):
    topic = guess_topic(rel)
    try:
        raw = read_docx(path) if path.lower().endswith(".docx") else read_pdf(path)
    except Exception as e:
        skipped.append((rel, "打不开:" + type(e).__name__))
        continue

    if is_garbled(raw):
        skipped.append((rel, "乱码/无文字层"))
        continue

    text = clean_text(raw)
    if len(text) < MIN_CHARS:
        skipped.append((rel, "内容过短(%d字)" % len(text)))
        continue

    # 按【内容指纹】去重：文件名不可靠（老师给的名字本身就乱）
    key = sig(text)
    if key in seen_names:
        prev = seen_names[key]
        if prev["chars"] >= len(text):
            skipped.append((rel, "内容重复（已有更全版本：%s）" % prev["source"][-30:]))
            continue
        old_file = prev["file"]
        manifest = [m for m in manifest if m is not prev]
        try:
            os.remove(os.path.join(OUT_DIR, old_file))
        except OSError:
            pass

    out_name = "%03d_%s.txt" % (i, safe_name(rel))
    out_path = os.path.join(OUT_DIR, out_name)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("【主题】%s\n【来源】%s\n\n%s\n" % (topic, rel, text))

    entry = {
        "index": i,
        "source": rel,
        "topic": topic,
        "chars": len(text),
        "file": out_name,
    }
    manifest.append(entry)
    seen_names[key] = entry

with open(MANIFEST, "w", encoding="utf-8") as f:
    json.dump(manifest, f, ensure_ascii=False, indent=2)

print("=" * 66)
print("成功提取：%d 个文档" % len(manifest))
print("跳过：%d 个" % len(skipped))
for rel, why in skipped:
    print("   -", rel, "->", why)
print()
print("按主题统计：")
by_topic = {}
for m in manifest:
    by_topic.setdefault(m["topic"], []).append(m)
for topic, items in sorted(by_topic.items(), key=lambda x: -sum(i["chars"] for i in x[1])):
    total = sum(i["chars"] for i in items)
    print("   %-14s %2d 篇  %6d 字" % (topic, len(items), total))
print()
print("总计：%d 字" % sum(m["chars"] for m in manifest))
print("输出目录：", OUT_DIR)